# -*- coding: utf-8 -*-
"""
S0_35_Elasticity_signs_revise.py

Correct cross-price elasticity signs using a sign matrix; modify only two sheets while exporting all sheets.

User-specified interpretation:
- Row Commodity represents quantity change Delta_Q_x, in demand or supply of commodity x.
- Columns Y1,Y2,... represent cross-commodity price changes Delta_P_y.
Each cell is e_{x,y} = d ln Q_x / d ln P_y.

Input 1: sign matrix.
  ../../input/Driver/Elasticity/Item_demand_supplu_cross_prodchain_signs.xlsx
    - demand_cross_signs
    - supply_cross_signs
  Codes:
    0: No cross-elasticity relationship; set the corresponding cell to zero.
    1: Preserve the original value and sign.
    +: Force positive using abs(value), preserving magnitude.
    -: Force negative using -abs(value), preserving magnitude.
  Force diagonal entries x==y to zero.

Input 2: original elasticity workbook read from retired_raw_unused.
  ../../input/Driver/Elasticity/retired_raw_unused/Elasticity_v3_processed_filled_by_region.xlsx
  Only two sheets are modified:
    - Demand_Cross_mean
    - Supply_Cross_mean

Output the full workbook, preserving all sheets except the two corrected tables.
  ../../input/Driver/Elasticity/Elasticity_v3_processed_filled_by_region.xlsx
"""

from __future__ import annotations
import argparse
from pathlib import Path
from typing import Dict, Optional, Any, Tuple, List

import pandas as pd
import openpyxl


TARGET_SHEETS = {
    "Demand_Cross_mean": "demand_cross_signs",
    "Supply_Cross_mean": "supply_cross_signs",
}


def _resolve(p: str | Path, base: Path) -> Path:
    p = Path(p)
    return p if p.is_absolute() else (base / p).resolve()


def read_sign_matrix(signs_xlsx: Path, sheet: str) -> pd.DataFrame:
    """Read a commodity-by-commodity sign matrix; normalize values to '0','1','+','-', or None."""
    df = pd.read_excel(signs_xlsx, sheet_name=sheet, dtype=object)
    if df.shape[1] < 2:
        raise ValueError(f"Sign sheet '{sheet}' empty/malformed: {signs_xlsx}")

    # The first column contains row commodity names.
    row_key = df.columns[0]
    df = df.rename(columns={row_key: "Commodity"})
    df["Commodity"] = df["Commodity"].astype(str).str.strip()
    df = df.set_index("Commodity")
    df.columns = [str(c).strip() for c in df.columns]

    def norm(v: Any) -> Optional[str]:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return None
        s = str(v).strip()
        if s == "":
            return None
        if s in {"+", "-"}:
            return s
        if s in {"0", "1"}:
            return s
        # Excel may store 0/1 as numeric values.
        try:
            fv = float(s)
            if fv == 0:
                return "0"
            if fv == 1:
                return "1"
        except Exception:
            pass
        return None  # Unrecognized codes mean no rule; leave unchanged.

    return df.applymap(norm)


def build_sign_dict(sign_mat: pd.DataFrame) -> Dict[str, Dict[str, str]]:
    """Convert the small matrix to sign_dict[x][y] -> code for fast lookup."""
    out: Dict[str, Dict[str, str]] = {}
    for x in sign_mat.index:
        row = sign_mat.loc[x]
        d = {}
        for y, code in row.items():
            if code is None:
                continue
            d[str(y)] = str(code)
        out[str(x)] = d
    return out


def find_header_row_and_col_map(ws) -> Tuple[int, Dict[str, int], int]:
    """Find the header row containing Commodity and build column-index mappings."""
    max_scan = min(ws.max_row, 30)
    header_row = None
    commodity_col = None

    for r in range(1, max_scan + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            if str(v).strip().lower() == "commodity":
                header_row = r
                commodity_col = c
                break
        if header_row is not None:
            break

    if header_row is None or commodity_col is None:
        raise ValueError(f"Cannot find header row with 'Commodity' in sheet '{ws.title}'")

    col_map: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        name = ws.cell(header_row, c).value
        if name is None:
            continue
        s = str(name).strip()
        if s == "":
            continue
        col_map[s] = c

    return header_row, col_map, commodity_col


def _to_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s == "":
        return None
    try:
        return float(s)
    except Exception:
        return None


def revise_sheet_inplace(ws, sign_mat: pd.DataFrame) -> Dict[str, int]:
    """
    Modify ws in place using sign_mat, only for Commodity-by-commodity cross-elasticity cells.
    Return summary counts.
    """
    header_row, col_map, commodity_col = find_header_row_and_col_map(ws)

    # Price-commodity columns matching sign-matrix columns
    price_cols: List[str] = [name for name in col_map.keys() if name in sign_mat.columns]
    if not price_cols:
        raise ValueError(
            f"Sheet '{ws.title}': no header columns match sign matrix columns. "
            f"Example headers: {list(col_map.keys())[:10]}"
        )

    # Convert to a dictionary for speed.
    sign_mat_sub = sign_mat[price_cols]
    sign_dict = build_sign_dict(sign_mat_sub)

    price_info = [(y, col_map[y]) for y in price_cols]
    max_col = max([commodity_col] + [ci for _, ci in price_info])

    counts = dict(touched=0, forced_pos=0, forced_neg=0, set_zero=0, diag_zero=0, skipped_missing_map=0)

    # Use iter_rows for speed; min_col=1 allows cell access with col_idx-1.
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, min_col=1, max_col=max_col):
        x_cell = row[commodity_col - 1]
        if x_cell.value is None:
            continue
        x = str(x_cell.value).strip()
        if x == "" or x.lower() == "nan":
            continue

        row_signs = sign_dict.get(x)
        if row_signs is None:
            continue

        for y, col_idx in price_info:
            cell = row[col_idx - 1]

            # Force the diagonal to zero.
            if y == x:
                v0 = _to_float(cell.value)
                if v0 is None or v0 != 0.0:
                    cell.value = 0.0
                    counts["diag_zero"] += 1
                continue

            code = row_signs.get(y)
            if code is None:
                counts["skipped_missing_map"] += 1
                continue

            if code == "1":
                continue

            v = _to_float(cell.value)
            if v is None:
                continue

            if code == "0":
                if v != 0.0:
                    cell.value = 0.0
                    counts["set_zero"] += 1
                    counts["touched"] += 1
                continue

            if code == "+":
                nv = abs(v)
                if nv != v:
                    cell.value = nv
                    counts["forced_pos"] += 1
                    counts["touched"] += 1
                continue

            if code == "-":
                nv = -abs(v)
                if nv != v:
                    cell.value = nv
                    counts["forced_neg"] += 1
                    counts["touched"] += 1
                continue

    return counts


def main():
    here = Path(__file__).resolve().parent

    parser = argparse.ArgumentParser(description="Revise cross-price elasticity signs using a sign matrix workbook.")
    parser.add_argument(
        "--signs",
        default="../../input/Driver/Elasticity/Item_demand_supplu_cross_prodchain_signs.xlsx",
        help="Signs workbook path",
    )
    parser.add_argument(
        "--elasticity",
        default="../../input/Driver/Elasticity/retired_raw_unused/Elasticity_v3_processed_filled_by_region.xlsx",
        help="Source elasticity workbook path (read from retired_raw_unused)",
    )
    parser.add_argument(
        "--out",
        default="../../input/Driver/Elasticity/Elasticity_v3_processed_filled_by_region.xlsx",
        help="Output path (FULL workbook). Only two target sheets are modified.",
    )
    args = parser.parse_args()

    signs_path = _resolve(args.signs, here)
    elasticity_path = _resolve(args.elasticity, here)
    out_path = _resolve(args.out, here)

    if not signs_path.exists():
        raise FileNotFoundError(f"Signs workbook not found: {signs_path}")
    if not elasticity_path.exists():
        raise FileNotFoundError(f"Elasticity workbook not found: {elasticity_path}")

    # Read the sign matrix.
    demand_sign = read_sign_matrix(signs_path, TARGET_SHEETS["Demand_Cross_mean"])
    supply_sign = read_sign_matrix(signs_path, TARGET_SHEETS["Supply_Cross_mean"])

    # Load the complete workbook, preserving all sheets.
    wb = openpyxl.load_workbook(elasticity_path)

    summary = {}
    for target_sheet in TARGET_SHEETS.keys():
        if target_sheet not in wb.sheetnames:
            raise KeyError(f"Target sheet '{target_sheet}' not found in {elasticity_path}. Found: {wb.sheetnames}")

        ws = wb[target_sheet]
        sign_mat = demand_sign if target_sheet == "Demand_Cross_mean" else supply_sign
        summary[target_sheet] = revise_sheet_inplace(ws, sign_mat)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    print("DONE. Full workbook written (all sheets preserved; only 2 sheets revised):")
    print(f"  Output: {out_path}")
    for sh, c in summary.items():
        print(f"- {sh}: touched={c['touched']}, forced_pos={c['forced_pos']}, forced_neg={c['forced_neg']}, "
              f"set_zero={c['set_zero']}, diag_zero={c['diag_zero']}, skipped_missing_map={c['skipped_missing_map']}")


if __name__ == "__main__":
    main()
